"""Краткий отчёт (справка) об аттестации кафедры — отдельно по бакалавриату и магистратуре:
сводные таблицы по группам, преподавателям и критическим дисциплинам на 3–5 страницах и
приложение А с баллами каждого студента по группам.

Реквизиты — по ГОСТ Р 7.0.97-2016 (справка), текст и таблицы — по ГОСТ 7.32-2017:
поля 30/15/20/20 мм, Times New Roman 14 пт, интервал 1,5, абзацный отступ 1,25 см,
«Таблица N – Название» над таблицей, повтор шапки таблицы на следующей странице.
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

import build_attestation_report as detailed
import report_model as rm


UNIVERSITY_LINES = [
    "Министерство науки и высшего образования Российской Федерации",
    "Федеральное государственное бюджетное образовательное учреждение высшего образования",
    "«Поволжский государственный технологический университет»",
]
CITY = "г. Йошкар-Ола"


def organization_lines(model: rm.DepartmentModel) -> list[str]:
    return UNIVERSITY_LINES + [model.department.faculty_title, f"Кафедра {model.department.name}"]

TEXT_SIZE = 14
TABLE_SIZE = 11  # ГОСТ 7.32 допускает в таблицах шрифт меньше, чем в тексте
HEADER_SIZE = 10
CONTENT_WIDTH_MM = 165  # 210 - 30 - 15
STATUS_FILLS = {rm.STATUS_OK: "E2EFDA", rm.STATUS_WARN: "FFF2CC", rm.STATUS_CRIT: "F8D7DA"}
SUBHEADER_FILL = "F2F2F2"
# баллы в приложении: ниже минимального, 0 баллов, у всей группы 0 (баллы не выставлены)
SCORE_FILLS = {"below": "FFF2CC", "zero": "F8D7DA", "none": "E7E6E6"}
REPORT_NAMES = {
    "б": "attestation_report_short_bachelor.docx",
    "м": "attestation_report_short_master.docx",
}


class Numbering:
    def __init__(self) -> None:
        self.table = 0
        self.section = 0

    def next_table(self) -> int:
        self.table += 1
        return self.table

    def next_section(self) -> int:
        self.section += 1
        return self.section


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(TEXT_SIZE)
    normal.font.color.rgb = detailed.BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detailed.add_page_number_field(footer_paragraph)


def add_paragraph(
    container,
    text: str = "",
    *,
    align: int = WD_ALIGN_PARAGRAPH.JUSTIFY,
    indent: bool = True,
    bold: bool = False,
    size: int = TEXT_SIZE,
    spacing: float = 1.5,
    space_before: int = 0,
    space_after: int = 0,
    keep_with_next: bool = False,
):
    paragraph = container.add_paragraph()
    paragraph.alignment = align
    paragraph_format = paragraph.paragraph_format
    paragraph_format.first_line_indent = Mm(12.5) if indent else Mm(0)
    paragraph_format.line_spacing = spacing
    paragraph_format.space_before = Pt(space_before)
    paragraph_format.space_after = Pt(space_after)
    paragraph_format.keep_with_next = keep_with_next
    if text:
        run = paragraph.add_run(text)
        detailed.apply_run_font(run, size=size, bold=bold)
    return paragraph


def add_heading(document: Document, numbering: Numbering, text: str) -> None:
    """Заголовок раздела по ГОСТ 7.32: номер, с абзацного отступа, полужирный, без точки в конце."""
    add_paragraph(document, f"{numbering.next_section()} {text}", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, space_before=12, space_after=6, keep_with_next=True)


def add_table_caption(document: Document, numbering: Numbering, title: str, number: str | None = None) -> None:
    """«Таблица 2 – …»; в приложении номер свой: «Таблица А.1 – …»."""
    add_paragraph(
        document,
        f"Таблица {number or numbering.next_table()} – {title}",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        indent=False,
        spacing=1.0,
        space_before=6,
        space_after=3,
        keep_with_next=True,
    )


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def mark_header_row(row) -> None:
    """Повтор строки на каждой странице и запрет разрыва строки."""
    properties = row._tr.get_or_add_trPr()
    for tag in ("w:tblHeader", "w:cantSplit"):
        element = OxmlElement(tag)
        element.set(qn("w:val"), "true")
        properties.append(element)


def forbid_row_split(row) -> None:
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(element)


def set_cell(
    cell,
    text: str,
    *,
    align: int = WD_ALIGN_PARAGRAPH.CENTER,
    bold: bool = False,
    fill: str | None = None,
    size: int = TABLE_SIZE,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Mm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    detailed.apply_run_font(run, size=size, bold=bold)
    if fill:
        shade(cell, fill)


def add_grid_table(document: Document, headers: list[str], widths_mm: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    fix_column_widths(table, widths_mm)
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, bold=True, size=HEADER_SIZE)
    mark_header_row(table.rows[0])
    return table


def fix_column_widths(table, widths_mm: list[float]) -> None:
    """Word читает ширину из ячеек, LibreOffice — из сетки таблицы; задаём и то и другое."""
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for column, width in zip(table.columns, widths_mm):
        column.width = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)


def add_row(
    table,
    values: list[str],
    widths_mm: list[float],
    *,
    left_columns: set[int] = frozenset(),
    bold: bool | set[int] = False,
    fills: dict[int, str] | None = None,
    size: int = TABLE_SIZE,
):
    """bold — вся строка или только перечисленные столбцы."""
    row = table.add_row()
    forbid_row_split(row)
    for index, (cell, value) in enumerate(zip(row.cells, values)):
        cell.width = Mm(widths_mm[index])
        align = WD_ALIGN_PARAGRAPH.LEFT if index in left_columns else WD_ALIGN_PARAGRAPH.CENTER
        heavy = index in bold if isinstance(bold, set) else bold
        set_cell(cell, value, align=align, bold=heavy, fill=(fills or {}).get(index), size=size)
    return row


def add_subheader_row(table, text: str) -> None:
    row = table.add_row()
    forbid_row_split(row)
    merged = row.cells[0].merge(row.cells[-1])
    set_cell(merged, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, fill=SUBHEADER_FILL)


def add_note(document: Document, text: str) -> None:
    add_paragraph(document, f"Примечание – {text}", size=12, spacing=1.0, space_before=3, space_after=6)


def model_levels(model: rm.DepartmentModel) -> list[str]:
    return [level for level in rm.LEVEL_TITLES if model.groups_of(level)]


def add_requisites(document: Document, model: rm.DepartmentModel) -> None:
    for line in organization_lines(model):
        add_paragraph(document, line, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12, spacing=1.0)
    add_paragraph(document, "СПРАВКА", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True, space_before=18, space_after=6)

    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    fix_column_widths(table, [CONTENT_WIDTH_MM / 2, CONTENT_WIDTH_MM / 2])
    set_cell(table.rows[0].cells[0], f"{datetime.now():%d.%m.%Y} №{rm.NBSP}______", align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    set_cell(table.rows[0].cells[1], CITY, align=WD_ALIGN_PARAGRAPH.RIGHT, size=12)

    levels = model_levels(model)
    whose = f"студентов {rm.LEVEL_GENITIVE[levels[0]]}" if len(levels) == 1 else "студентов"
    add_paragraph(
        document,
        f"О результатах аттестации № {model.attestation} {whose} в {model.semester.replace('ый', 'ом')} "
        f"семестре {model.study_year} учебного года",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        indent=False,
        bold=True,
        spacing=1.0,
        space_before=12,
        space_after=12,
    )


def add_overview(document: Document, model: rm.DepartmentModel) -> None:
    levels = model_levels(model)
    if len(levels) == 1:
        groups = f"{len(model.groups)} учебных групп {rm.LEVEL_GENITIVE[levels[0]]} кафедры"
    else:
        counts = ", ".join(f"{rm.LEVEL_TITLES[level]} — {len(model.groups_of(level))}" for level in levels)
        groups = f"{len(model.groups)} учебных групп кафедры ({counts})"
    add_paragraph(
        document,
        f"По данным системы РИТМ на {model.execution_time} в аттестации учтены {model.students} студентов "
        f"{groups}. По всем дисциплинам аттестованы "
        f"{model.fully_attested} студентов ({rm.format_percent(rm.rate(model.fully_attested, model.students))}). "
        f"Из {model.records} оценок по дисциплинам не аттестовано {model.failed} "
        f"({rm.format_percent(model.failure_rate)}), в том числе с нулевым баллом — {model.zero}.",
    )
    if model.rank is not None:
        add_paragraph(
            document,
            f"По доле неаттестаций кафедра {model.department.name} занимает {model.rank.place}-е место из "
            f"{model.rank.total} кафедр университета: {rm.format_percent(model.rank.failure_rate)} "
            f"при среднем значении {rm.format_percent(model.rank.university_rate)} "
            f"(по всем дисциплинам, которые ведут преподаватели кафедры).",
        )


def add_groups_table(document: Document, model: rm.DepartmentModel, numbering: Numbering) -> None:
    add_heading(document, numbering, "Результаты аттестации по группам")
    has_delta = any(group.delta is not None for group in model.groups)
    headers = ["Группа", "Курс", "Студентов", "Аттестованы полностью, чел.", "Н/А, %"]
    widths = [32, 13, 24, 32, 22]
    if has_delta:
        headers.append("Изменение, п. п.")
        widths.append(18)
    headers.append("Статус")
    widths.append(CONTENT_WIDTH_MM - sum(widths))

    add_table_caption(document, numbering, "Результаты аттестации по группам")
    table = add_grid_table(document, headers, widths)

    def values(name: str, course: str, groups: list[rm.GroupSummary], delta: float | None) -> tuple[list[str], str]:
        students = sum(len(group.students) for group in groups)
        records = sum(group.records for group in groups)
        failed = sum(group.failed for group in groups)
        fully = sum(group.fully_attested for group in groups)
        failure_rate = rm.rate(failed, records)
        status = rm.status_for(failure_rate)
        row = [name, course, str(students), f"{fully} ({rm.format_percent(rm.rate(fully, students))})", rm.format_percent(failure_rate)]
        if has_delta:
            row.append(rm.format_delta(delta))
        row.append(rm.status_label(status))
        return row, status

    status_columns = (4, len(headers) - 1)
    levels = model_levels(model)
    for level in levels:
        title = rm.LEVEL_TITLES[level]
        groups = model.groups_of(level)
        if len(levels) > 1:
            add_subheader_row(table, title.capitalize())
        for group in groups:
            row, status = values(group.name, str(group.course or "–"), [group], group.delta)
            add_row(table, row, widths, left_columns={0}, fills={column: STATUS_FILLS[status] for column in status_columns})
        row, status = values(f"Итого: {title}", "", groups, None)
        add_row(table, row, widths, left_columns={0}, bold=True, fills={column: STATUS_FILLS[status] for column in status_columns})
    if len(levels) > 1:
        row, status = values("Итого по кафедре", "", model.groups, None)
        add_row(table, row, widths, left_columns={0}, bold=True, fills={column: STATUS_FILLS[status] for column in status_columns})

    add_note(
        document,
        f"Н/А — неаттестация: балл ниже минимального по технологической карте. "
        f"Статус по доле неаттестаций: {rm.status_label(rm.STATUS_OK)} — не более {rm.WARN_THRESHOLD:.0f} %; "
        f"{rm.status_label(rm.STATUS_WARN)} — от {rm.WARN_THRESHOLD:.0f} до {rm.CRIT_THRESHOLD:.0f} %; "
        f"{rm.status_label(rm.STATUS_CRIT)} — более {rm.CRIT_THRESHOLD:.0f} %."
        + (" Изменение — разность с предыдущей аттестацией в процентных пунктах." if has_delta else ""),
    )
    add_paragraph(document, "Баллы каждого студента по дисциплинам приведены в приложении А.")


def add_teachers_table(document: Document, model: rm.DepartmentModel, numbering: Numbering) -> None:
    add_heading(document, numbering, "Результаты аттестации по преподавателям")
    add_paragraph(
        document,
        f"В таблицу включены все преподаватели, ведущие дисциплины в группах кафедры; преподаватели "
        f"сопоставлены с дисциплинами по отчету «Процент аттестаций по кафедрам» на {model.teachers_source_time}.",
    )
    headers = ["Преподаватель", "Кафедра", "Групп", "Дисциплин", "Оценок", "Н/А, %", "Статус"]
    widths = [38, 27, 14, 22, 18, 20, 26]
    add_table_caption(document, numbering, "Результаты аттестации по преподавателям")
    table = add_grid_table(document, headers, widths)
    for teacher in model.teachers:
        fill = STATUS_FILLS[teacher.status]
        add_row(
            table,
            [
                rm.short_name(teacher.name),
                teacher.kafedra,
                str(len(teacher.groups)),
                str(len(teacher.subjects)),
                str(teacher.records),
                rm.format_percent(teacher.failure_rate),
                rm.status_label(teacher.status),
            ],
            widths,
            left_columns={0, 1},
            fills={5: fill, 6: fill},
        )
    unassigned = [item for item in model.disciplines if not item.teachers]
    if unassigned:
        levels = {item.level for item in unassigned}
        reason = (
            "в отчете «Процент аттестаций по кафедрам» нет групп магистратуры"
            if levels == {"м"}
            else "дисциплины не найдены в отчете «Процент аттестаций по кафедрам» (в основном группы магистратуры и элективные дисциплины)"
        )
        add_note(document, f"Для {len(unassigned)} дисциплин в группах преподаватель не указан: {reason}.")


def add_critical_table(document: Document, model: rm.DepartmentModel, numbering: Numbering) -> None:
    add_heading(document, numbering, f"Дисциплины с долей неаттестаций {rm.CRITICAL_DISCIPLINE_THRESHOLD:.0f} % и более")
    critical = model.critical_disciplines()
    headers = ["Группа", "Дисциплина", "Преподаватель", "Н/А, чел.", "Н/А, %"]
    widths = [26, 68, 33, 20, 18]
    add_table_caption(document, numbering, f"Дисциплины с долей неаттестаций {rm.CRITICAL_DISCIPLINE_THRESHOLD:.0f} % и более")
    table = add_grid_table(document, headers, widths)
    if not critical:
        add_row(table, ["–", "Нет таких дисциплин", "–", "–", "–"], widths, left_columns={1})
    for item in critical:
        add_row(
            table,
            [item.group, item.subject, item.teacher_names, f"{item.failed} из {item.students}", rm.format_percent(item.failure_rate)],
            widths,
            left_columns={1, 2},
            fills={4: STATUS_FILLS[rm.STATUS_CRIT]},
        )


def add_conclusions(document: Document, model: rm.DepartmentModel, numbering: Numbering) -> None:
    critical_table = numbering.table  # таблица критических дисциплин — последняя перед выводами
    add_heading(document, numbering, "Выводы")
    for line in rm.conclusions(model, critical_reference=f"в таблице {critical_table}"):
        add_paragraph(document, f"– {line}")


def add_signature(document: Document, model: rm.DepartmentModel) -> None:
    add_paragraph(document, "", space_before=18)
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    fix_column_widths(table, [80, 40, 45])
    for cell, text, align in zip(
        table.rows[0].cells,
        [f"Заведующий кафедрой {model.department.name}", "____________", model.department.head_name],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT],
    ):
        set_cell(cell, text, align=align, size=TEXT_SIZE)


def student_labels(students: list[rm.StudentSummary]) -> dict[str, str]:
    """Фамилия и инициалы; если они совпали у двух студентов группы — полное ФИО."""
    short = {student.name: rm.short_name(student.name) for student in students}
    repeated = Counter(short.values())
    return {name: label if repeated[label] == 1 else name for name, label in short.items()}


def format_score(value: float) -> str:
    """Один знак после запятой: «94,91» не помещается в столбец балла."""
    return f"{round(value, 1):g}".replace(".", ",")


def narrow_cell_margins(table, mm: float) -> None:
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(round(mm * 56.7)))  # твипы: 1 мм = 56,7
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    table._tbl.tblPr.append(margins)


def score_fill(student: rm.StudentSummary, item: rm.DisciplineSummary) -> str | None:
    if item.not_entered:
        return SCORE_FILLS["none"]
    if item.subject not in student.failed_subjects:
        return None
    return SCORE_FILLS["zero"] if student.scores[item.subject] == 0 else SCORE_FILLS["below"]


def add_group_students(document: Document, numbering: Numbering, group: rm.GroupSummary, number: str) -> None:
    disciplines = group.disciplines  # порядок тот же, что на тепловой карте сайта: по доле Н/А
    widths = [8, 44]
    rest = CONTENT_WIDTH_MM - sum(widths) - 14
    widths += [rest / max(len(disciplines), 1)] * len(disciplines) + [14]
    size = 9 if len(disciplines) <= 10 else 8

    add_table_caption(document, numbering, f"Группа {group.name}", number=number)
    table = add_grid_table(document, ["№", "Студент", *(str(index) for index in range(1, len(disciplines) + 1)), "Н/А"], widths)
    narrow_cell_margins(table, 0.8)
    minimum = add_row(
        table,
        ["", "Минимальный балл", *(format_score(item.threshold) for item in disciplines), ""],
        widths,
        left_columns={1},
        fills={index: SUBHEADER_FILL for index in range(len(widths))},
        size=size,
    )
    mark_header_row(minimum)

    labels = student_labels(group.students)
    for index, student in enumerate(sorted(group.students, key=lambda item: item.name), start=1):
        cells = [str(index), labels[student.name]]
        fills = {}
        for column, item in enumerate(disciplines, start=2):
            score = student.scores.get(item.subject)
            cells.append("–" if score is None else format_score(score))
            fill = None if score is None else score_fill(student, item)
            if fill:
                fills[column] = fill
        cells.append(str(student.failed))
        add_row(
            table, cells, widths, left_columns={1}, fills=fills, size=size,
            bold={len(cells) - 1} if student.failed else False,
        )

    entries = []
    for index, item in enumerate(disciplines, start=1):
        teachers = ", ".join(rm.short_name(name) for name, _ in item.teachers)
        entries.append(f"{index} – {item.subject}" + (f" ({teachers})" if teachers else ""))
    add_paragraph(document, "Дисциплины: " + "; ".join(entries) + ".", size=10, spacing=1.0, space_before=3, space_after=9)


def add_students_appendix(document: Document, model: rm.DepartmentModel, numbering: Numbering) -> None:
    """Приложение по ГОСТ 7.32: с новой страницы, «ПРИЛОЖЕНИЕ А» и заголовок по центру."""
    title = add_paragraph(document, "ПРИЛОЖЕНИЕ А", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True)
    title.paragraph_format.page_break_before = True
    add_paragraph(
        document, "Результаты аттестации студентов по группам",
        align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True, space_after=6,
    )
    add_paragraph(
        document,
        "В таблицах — аттестационный балл каждого студента по каждой дисциплине группы; номера столбцов "
        "расшифрованы под таблицей, в строке «Минимальный балл» — порог аттестации. Заливкой отмечены: "
        "жёлтой — балл ниже минимального, розовой — 0 баллов, серой — дисциплина, по которой у всей группы "
        "0 баллов (баллы не выставлены); «–» — дисциплины нет в плане студента. Н/А — число неаттестаций студента.",
        size=12, spacing=1.0, space_after=6,
    )
    for index, group in enumerate(model.groups, start=1):
        add_group_students(document, numbering, group, f"А.{index}")


def build_short_report(attestation_dir: Path, level: str, output_path: Path | None = None) -> Path:
    """Справка по одному уровню образования: «б» — бакалавриат, «м» — магистратура."""
    model = rm.build_model(attestation_dir, level)
    document = Document()
    configure_document(document)
    numbering = Numbering()

    add_requisites(document, model)
    add_overview(document, model)
    add_groups_table(document, model, numbering)
    if model.teachers:
        add_teachers_table(document, model, numbering)
    add_critical_table(document, model, numbering)
    add_conclusions(document, model, numbering)
    add_signature(document, model)
    add_students_appendix(document, model, numbering)

    output_path = output_path or attestation_dir / REPORT_NAMES[level]
    document.save(str(output_path))
    return output_path


def build_short_reports(attestation_dir: Path) -> list[Path]:
    """Справки по каждому уровню, для которого выгружены группы."""
    return [
        build_short_report(attestation_dir, level)
        for level in REPORT_NAMES
        if any((attestation_dir / level).glob("*.xml"))
    ]


def main() -> int:
    parser = argparse.ArgumentParser(description="Краткий отчёт (справка) об аттестации кафедры.")
    parser.add_argument("attestation_dir", help="Папка аттестации, например 2026-2027/нечет/1")
    args = parser.parse_args()
    try:
        paths = build_short_reports(Path(args.attestation_dir).resolve())
    except Exception as exc:  # pragma: no cover
        print(f"Ошибка сборки краткого отчёта: {exc}", file=sys.stderr)
        return 1
    for path in paths:
        print(f"Краткий отчёт готов: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
