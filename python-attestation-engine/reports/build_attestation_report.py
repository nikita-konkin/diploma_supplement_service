from __future__ import annotations

import argparse
import re
import sys
import xml.etree.ElementTree as ET
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

import kafedra_teachers
from department_info import load_department


BLACK = RGBColor(0, 0, 0)
IMAGE_WIDTH_MM = 155
FAILURE_HIGHLIGHT = 85.0
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
GENERAL_IMAGE_CAPTIONS = {
    "attestation_progress_comparison": "Динамика количества аттестованных записей по группам между соседними аттестациями.",
    "group_comparison_boxplot": "Сравнение групп по распределению результатов.",
    "group_risk_index": "Индекс риска учебных групп.",
    "top_problematic_subjects": "Наиболее проблемные дисциплины.",
}
GROUP_IMAGE_CAPTION_BUILDERS = {
    "by_stud_": lambda group_name: f"Распределение результатов студентов группы {group_name}.",
    "by_sub_": lambda group_name: f"Распределение результатов по дисциплинам группы {group_name}.",
}


@dataclass
class SubjectResult:
    name: str
    threshold: float
    att_ball: float
    is_attested: bool


@dataclass
class StudentRecord:
    full_name: str
    student_number: str
    subjects: list[SubjectResult]
    failed_count: int
    attested_count: int

    @property
    def failure_rate(self) -> float:
        total = len(self.subjects)
        return self.failed_count / total * 100 if total else 0.0


@dataclass
class SubjectAggregate:
    name: str
    total_students: int
    failed_students: int
    attested_students: int
    failure_rate: float


@dataclass
class GroupReport:
    name: str
    report_title: str
    execution_time: str
    students: list[StudentRecord]
    subject_stats: list[SubjectAggregate]
    total_records: int
    total_failed: int
    total_attested: int
    failure_rate: float
    source_path: Path


@dataclass
class AttestationComparison:
    current_attestation: int
    previous_attestation: int
    previous_dir: Path
    image_path: Path
    current_total_attested: int
    previous_total_attested: int
    rows: list["ComparisonRow"]


@dataclass
class ComparisonRow:
    group_name: str
    previous_attested: int
    current_attested: int
    delta_attested: int
    delta_attended_pct: float | None
    previous_failed: int
    current_failed: int
    student_count: int
    critical_note_lines: list[str]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a DOCX attestation report for bachelor study groups."
    )
    parser.add_argument(
        "--input-dir",
        default=None,
        help="Folder with XML source files, summary text report and charts.",
    )
    parser.add_argument(
        "--output",
        default="attestation_report_bachelor.docx",
        help="Path to the resulting DOCX report. Relative paths are resolved inside input-dir.",
    )
    return parser.parse_args()


def console_safe(text: str) -> str:
    encoding = sys.stdout.encoding or "utf-8"
    return text.encode(encoding, errors="backslashreplace").decode(encoding)


def prompt_for_input_dir() -> Path:
    default_dir = Path.cwd().resolve()
    prompt = (
        "Enter path to the folder with attestation data.\n"
        f"Press Enter to use the current directory: {default_dir}\n"
        "InputDir: "
    )

    while True:
        try:
            sys.stdout.write(console_safe(prompt))
            sys.stdout.flush()
            raw_value = input()
        except EOFError as exc:
            raise ValueError(
                "Input directory was not provided and interactive input is unavailable."
            ) from exc

        candidate_text = raw_value.strip().strip('"')
        candidate = Path(candidate_text) if candidate_text else default_dir
        candidate = candidate.expanduser()

        try:
            resolved = candidate.resolve()
        except OSError:
            print(console_safe("Invalid path. Try again."))
            continue

        if not resolved.exists():
            print(console_safe(f"Path does not exist: {resolved}"))
            continue
        if not resolved.is_dir():
            print(console_safe(f"Path is not a directory: {resolved}"))
            continue
        return resolved


def parse_threshold(value: str) -> float:
    match = re.search(r"[-+]?\d+(?:[.,]\d+)?", value or "")
    if not match:
        return 0.0
    return float(match.group(0).replace(",", "."))


def parse_score(value: str | None) -> float:
    if not value:
        return 0.0
    return float(value.replace(",", "."))


def nsmap(root: ET.Element) -> dict[str, str]:
    namespace = ""
    if root.tag.startswith("{"):
        namespace = root.tag[1:].split("}", 1)[0]
    return {"r": namespace}


def group_name_from_title(title: str, fallback: str) -> str:
    match = re.search(r"группы\s+(.+)$", title)
    return match.group(1).strip() if match else fallback


def parse_group(xml_path: Path) -> GroupReport:
    root = ET.parse(xml_path).getroot()
    ns = nsmap(root)
    title = root.attrib.get("ReportTitle", xml_path.stem)
    group_name = group_name_from_title(title, xml_path.stem)
    execution_time = root.attrib.get("ExecutionTime", "").strip()

    tablix = root.find("r:Tablix1", ns)
    if tablix is None:
        raise ValueError(f"Не найден узел Tablix1 в файле {xml_path}")

    student_collection = tablix.find("r:stFIO_Collection", ns)
    if student_collection is None:
        raise ValueError(f"Не найден список студентов в файле {xml_path}")

    students: list[StudentRecord] = []
    subject_totals: dict[str, int] = defaultdict(int)
    subject_failed: dict[str, int] = defaultdict(int)

    for wrapper in student_collection.findall("r:stFIO", ns):
        textbox = wrapper.find("r:Textbox48", ns)
        if textbox is None:
            continue
        student_el = textbox.find("r:stFIO", ns)
        if student_el is None:
            continue

        full_name = student_el.attrib.get("stFIO", "").strip()
        student_number_el = student_el.find("r:StudentNumber", ns)
        if student_number_el is None:
            continue

        student_number = student_number_el.attrib.get("StudentNumber", "").strip()
        subject_collection = student_number_el.find("r:SubjectName_Collection", ns)
        if subject_collection is None:
            continue

        subjects: list[SubjectResult] = []
        for subject_el in subject_collection.findall("r:SubjectName", ns):
            subject_name = subject_el.attrib.get("SubjectName", "").strip()
            min_ball_el = subject_el.find("r:minBall", ns)
            threshold = 0.0
            att_ball = 0.0
            if min_ball_el is not None:
                threshold = parse_threshold(min_ball_el.attrib.get("minBall", "").strip())
                att_ball = parse_score(min_ball_el.attrib.get("AttBall"))

            is_attested = att_ball >= threshold
            subjects.append(
                SubjectResult(
                    name=subject_name,
                    threshold=threshold,
                    att_ball=att_ball,
                    is_attested=is_attested,
                )
            )
            subject_totals[subject_name] += 1
            if not is_attested:
                subject_failed[subject_name] += 1

        failed_count = sum(1 for subject in subjects if not subject.is_attested)
        attested_count = len(subjects) - failed_count
        students.append(
            StudentRecord(
                full_name=full_name,
                student_number=student_number,
                subjects=subjects,
                failed_count=failed_count,
                attested_count=attested_count,
            )
        )

    summary_subjects: dict[str, int] = {}
    summary_box = tablix.find("r:Textbox53", ns)
    if summary_box is not None:
        summary_collection = summary_box.find("r:SubjectName_Collection", ns)
        if summary_collection is not None:
            for subject_el in summary_collection.findall("r:SubjectName", ns):
                subject_name = subject_el.attrib.get("SubjectName", "").strip()
                min_ball_el = subject_el.find("r:minBall", ns)
                if min_ball_el is not None:
                    summary_subjects[subject_name] = int(
                        min_ball_el.attrib.get("Textbox55", "0")
                    )

    subject_stats: list[SubjectAggregate] = []
    for subject_name, total_students in sorted(subject_totals.items()):
        failed_students = summary_subjects.get(subject_name, subject_failed[subject_name])
        attested_students = total_students - failed_students
        failure_rate = failed_students / total_students * 100 if total_students else 0.0
        subject_stats.append(
            SubjectAggregate(
                name=subject_name,
                total_students=total_students,
                failed_students=failed_students,
                attested_students=attested_students,
                failure_rate=failure_rate,
            )
        )

    subject_stats.sort(key=lambda item: (-item.failure_rate, -item.failed_students, item.name))
    total_records = sum(len(student.subjects) for student in students)
    total_failed = sum(student.failed_count for student in students)
    total_attested = total_records - total_failed
    failure_rate = total_failed / total_records * 100 if total_records else 0.0

    return GroupReport(
        name=group_name,
        report_title=title,
        execution_time=execution_time,
        students=sorted(students, key=lambda item: (-item.failed_count, item.full_name)),
        subject_stats=subject_stats,
        total_records=total_records,
        total_failed=total_failed,
        total_attested=total_attested,
        failure_rate=failure_rate,
        source_path=xml_path,
    )


def parse_all_groups(input_dir: Path) -> list[GroupReport]:
    xml_files = sorted(input_dir.glob("*.xml"))
    if not xml_files:
        raise FileNotFoundError(f"В папке {input_dir} не найдены XML-файлы групп.")

    unique_groups: dict[str, GroupReport] = {}
    for path in xml_files:
        group = parse_group(path)
        existing = unique_groups.get(group.name)
        if existing is None or should_replace_group(existing, group):
            unique_groups[group.name] = group

    return list(unique_groups.values())


def should_replace_group(existing: GroupReport, candidate: GroupReport) -> bool:
    existing_exact = existing.source_path.stem == existing.name
    candidate_exact = candidate.source_path.stem == candidate.name
    if existing_exact != candidate_exact:
        return candidate_exact
    if existing.total_records != candidate.total_records:
        return candidate.total_records > existing.total_records
    return len(candidate.source_path.name) < len(existing.source_path.name)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.first_line_indent = Mm(12.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    for style_name, font_size in (("Title", 16), ("Heading 1", 16), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(font_size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.first_line_indent = Mm(0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.space_before = Pt(0)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number_field(footer_paragraph)


def add_page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    apply_run_font(run, 12)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def apply_run_font(run, size: int = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def add_text(
    container,
    text: str,
    *,
    style: str | None = None,
    align: int | None = None,
    bold: bool = False,
    italic: bool = False,
    size: int = 12,
    first_line_indent_mm: float | None = 12.5,
) -> None:
    paragraph = container.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    if first_line_indent_mm is not None:
        paragraph.paragraph_format.first_line_indent = Mm(first_line_indent_mm)
    run = paragraph.add_run(text)
    apply_run_font(run, size=size, bold=bold, italic=italic)


def add_bullet(container, text: str) -> None:
    add_text(container, f"• {text}", first_line_indent_mm=0)


def format_percent(value: float) -> str:
    return f"{value:.1f}%".replace(".", ",")


def calc_percent_change(previous_value: int, current_value: int) -> float | None:
    if previous_value == 0:
        return 0.0 if current_value == 0 else None
    return (current_value - previous_value) / previous_value * 100


def format_signed_number(value: int) -> str:
    return f"{value:+d}"


def format_signed_percent(value: float | None) -> str:
    if value is None:
        return "н/д"
    return f"{value:+.1f}%".replace(".", ",")


def format_unsigned_percent(value: float | None) -> str:
    if value is None:
        return "н/д"
    return f"{abs(value):.1f}%".replace(".", ",")


def format_delta_with_percent(delta_value: int, delta_percent: float | None) -> str:
    return f"{format_signed_number(delta_value)} ({format_signed_percent(delta_percent)})"


def format_magnitude_with_percent(delta_value: int, delta_percent: float | None) -> str:
    return f"{abs(delta_value)} ({format_unsigned_percent(delta_percent)})"


def set_cell_text(cell, text: str, *, bold: bool = False, align: int = WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Mm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    apply_run_font(run, size=12, bold=bold)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            align = WD_ALIGN_PARAGRAPH.CENTER if index else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[index], value, align=align)


def add_figure(document: Document, image_path: Path, caption_text: str, figure_number: int) -> int:
    if not image_path.exists():
        return figure_number

    document.add_picture(str(image_path), width=Mm(IMAGE_WIDTH_MM))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.first_line_indent = Mm(0)
    caption_paragraph.paragraph_format.line_spacing = 1.0
    caption_paragraph.paragraph_format.space_after = Pt(6)
    run = caption_paragraph.add_run(f"Рисунок {figure_number} — {caption_text}")
    apply_run_font(run, size=12, italic=True)
    return figure_number + 1


def overview_rows(groups: list[GroupReport]) -> list[list[str]]:
    rows: list[list[str]] = []
    for group in groups:
        rows.append(
            [
                group.name,
                str(len(group.students)),
                str(len(group.subject_stats)),
                str(group.total_records),
                str(group.total_attested),
                str(group.total_failed),
                format_percent(group.failure_rate),
            ]
        )
    return rows


def top_problem_groups(groups: list[GroupReport], limit: int = 3) -> list[GroupReport]:
    return sorted(groups, key=lambda item: (-item.failure_rate, item.name))[:limit]


def top_success_groups(groups: list[GroupReport], limit: int = 3) -> list[GroupReport]:
    return sorted(groups, key=lambda item: (item.failure_rate, item.name))[:limit]


def overall_critical_subjects(groups: list[GroupReport], limit: int = 10) -> list[str]:
    rows: list[tuple[float, int, str]] = []
    for group in groups:
        for subject in group.subject_stats:
            if subject.failure_rate >= FAILURE_HIGHLIGHT:
                rows.append(
                    (
                        subject.failure_rate,
                        subject.failed_students,
                        f"{subject.name} ({group.name}) — {format_percent(subject.failure_rate)} "
                        f"({subject.failed_students} из {subject.total_students})",
                    )
                )
    rows.sort(key=lambda item: (-item[0], -item[1], item[2]))
    return [text for _, _, text in rows[:limit]]


def top_success_subjects(group: GroupReport, limit: int = 5) -> list[SubjectAggregate]:
    return sorted(
        group.subject_stats,
        key=lambda item: (item.failure_rate, item.failed_students, item.name),
    )[:limit]


def high_failure_students(group: GroupReport, limit: int = 10) -> list[StudentRecord]:
    students = [student for student in group.students if student.failure_rate > 85.0]
    return students[:limit]


def detect_previous_attestation_dir(input_dir: Path) -> tuple[int, int, Path] | None:
    try:
        current_attestation = int(input_dir.parent.name)
    except ValueError:
        return None

    if current_attestation <= 1:
        return None

    previous_attestation = current_attestation - 1
    previous_dir = input_dir.parent.parent / str(previous_attestation) / input_dir.name
    if not previous_dir.exists() or not previous_dir.is_dir():
        return None

    return current_attestation, previous_attestation, previous_dir


FONT_DIRS = [
    Path(r"C:\Windows\Fonts"),
    Path("/System/Library/Fonts/Supplemental"),
    Path("/Library/Fonts"),
    Path("/usr/share/fonts/truetype/liberation"),  # fonts-liberation в Docker-образе
]


def try_load_font(font_name: str, size: int) -> ImageFont.FreeTypeFont | None:
    for font_dir in FONT_DIRS:
        candidate = font_dir / font_name
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return None


def load_chart_font(size: int, *, bold: bool = False) -> ImageFont.ImageFont:
    candidates = (
        ["timesbd.ttf", "times.ttf", "timesbi.ttf", "timesi.ttf", "Times New Roman Bold.ttf", "Times New Roman.ttf",
         "LiberationSerif-Bold.ttf"]
        if bold
        else ["times.ttf", "timesi.ttf", "arial.ttf", "segoeui.ttf", "Times New Roman.ttf", "Arial.ttf",
              "LiberationSerif-Regular.ttf"]
    )
    for font_name in candidates:
        font = try_load_font(font_name, size)
        if font is not None:
            return font
    return ImageFont.load_default()


def shorten_subject_name(subject_name: str, max_length: int = 64) -> str:
    compact = re.sub(r"\s+", " ", subject_name).strip()
    if len(compact) <= max_length:
        return compact
    return compact[: max_length - 1].rstrip() + "…"


def build_critical_subject_note_lines(
    current_group: GroupReport | None,
    previous_group: GroupReport | None,
) -> list[str]:
    if current_group is None or previous_group is None:
        return []
    if current_group.total_failed <= previous_group.total_failed:
        return []
    if current_group.total_failed < len(current_group.students):
        return []

    critical_subjects = [
        subject
        for subject in current_group.subject_stats
        if subject.failure_rate > 95.0
    ]
    if not critical_subjects:
        return []

    failed_delta_pct = calc_percent_change(
        previous_group.total_failed,
        current_group.total_failed,
    )
    subject_bits = [
        f"{shorten_subject_name(subject.name)} — {format_percent(subject.failure_rate)}"
        for subject in critical_subjects[:2]
    ]
    if len(critical_subjects) > 2:
        subject_bits.append("и др.")

    lines = wrap(
        (
            "Рост неаттестаций: "
            f"{previous_group.total_failed} -> {current_group.total_failed} "
            f"({format_signed_percent(failed_delta_pct)}); "
            f"студентов в группе: {len(current_group.students)}."
        ),
        width=62,
    )
    note = "Критическая дисциплина: " + "; ".join(subject_bits)
    lines.extend(wrap(note, width=62))
    return lines


def build_attestation_comparison(input_dir: Path, current_groups: list[GroupReport]) -> AttestationComparison | None:
    attestation_context = detect_previous_attestation_dir(input_dir)
    if attestation_context is None:
        return None

    current_attestation, previous_attestation, previous_dir = attestation_context
    previous_groups = parse_all_groups(previous_dir)
    current_map = {group.name: group for group in current_groups}
    previous_map = {group.name: group for group in previous_groups}
    group_names = sorted(set(current_map) | set(previous_map))

    if not group_names:
        return None

    rows: list[ComparisonRow] = []
    for group_name in group_names:
        previous_group = previous_map.get(group_name)
        current_group = current_map.get(group_name)
        previous_value = previous_group.total_attested if previous_group is not None else 0
        current_value = current_group.total_attested if current_group is not None else 0
        delta_value = current_value - previous_value
        rows.append(
            ComparisonRow(
                group_name=group_name,
                previous_attested=previous_value,
                current_attested=current_value,
                delta_attested=delta_value,
                delta_attended_pct=calc_percent_change(previous_value, current_value),
                previous_failed=previous_group.total_failed if previous_group is not None else 0,
                current_failed=current_group.total_failed if current_group is not None else 0,
                student_count=len(current_group.students) if current_group is not None else (len(previous_group.students) if previous_group is not None else 0),
                critical_note_lines=build_critical_subject_note_lines(current_group, previous_group),
            )
        )

    image_path = input_dir / "attestation_progress_comparison.png"
    render_attestation_comparison_chart(
        image_path=image_path,
        rows=rows,
        current_attestation=current_attestation,
        previous_attestation=previous_attestation,
    )

    return AttestationComparison(
        current_attestation=current_attestation,
        previous_attestation=previous_attestation,
        previous_dir=previous_dir,
        image_path=image_path,
        current_total_attested=sum(row.current_attested for row in rows),
        previous_total_attested=sum(row.previous_attested for row in rows),
        rows=rows,
    )


def render_attestation_comparison_chart(
    image_path: Path,
    rows: list[ComparisonRow],
    current_attestation: int,
    previous_attestation: int,
) -> None:
    width = 1700
    top_margin = 210
    bottom_margin = 90
    left_margin = 330
    right_margin = 300
    plot_width = width - left_margin - right_margin
    center_x = left_margin + plot_width // 2
    base_row_height = 92
    note_line_height = 24
    row_heights = [
        base_row_height + note_line_height * len(row.critical_note_lines)
        for row in rows
    ]
    height = top_margin + bottom_margin + sum(row_heights)

    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    title_font = load_chart_font(36, bold=True)
    subtitle_font = load_chart_font(22)
    label_font = load_chart_font(24, bold=True)
    value_font = load_chart_font(22)
    legend_font = load_chart_font(20)
    note_font = load_chart_font(18)

    title = f"Динамика количества аттестованных записей по группам: аттестация №{current_attestation} к №{previous_attestation}"
    subtitle = (
        "Зеленый цвет показывает рост, красный — снижение, серый — отсутствие "
        "изменений; в скобках указан относительный процент изменения."
    )
    draw.text((50, 36), title, fill="black", font=title_font)
    draw.text((50, 92), subtitle, fill="#444444", font=subtitle_font)

    legend_y = 138
    draw.rectangle((50, legend_y, 82, legend_y + 20), fill="#2e8b57")
    draw.text((94, legend_y - 5), "Рост", fill="black", font=legend_font)
    draw.rectangle((180, legend_y, 212, legend_y + 20), fill="#c0392b")
    draw.text((224, legend_y - 5), "Снижение", fill="black", font=legend_font)
    draw.rectangle((380, legend_y, 412, legend_y + 20), fill="#8c8c8c")
    draw.text((424, legend_y - 5), "Без изменений", fill="black", font=legend_font)

    max_abs_delta = max(abs(row.delta_attested) for row in rows) if rows else 1
    max_abs_delta = max(max_abs_delta, 1)
    scale = (plot_width / 2 - 40) / max_abs_delta

    draw.line((center_x, top_margin - 20, center_x, height - bottom_margin + 20), fill="#404040", width=3)

    for tick in range(-max_abs_delta, max_abs_delta + 1):
        if tick == 0 or max_abs_delta > 12 and tick % max(1, max_abs_delta // 6) != 0:
            continue
        x = center_x + tick * scale
        draw.line((x, top_margin - 12, x, height - bottom_margin + 8), fill="#e0e0e0", width=1)
        text = f"{tick:+d}"
        bbox = draw.textbbox((0, 0), text, font=legend_font)
        draw.text((x - (bbox[2] - bbox[0]) / 2, top_margin - 46), text, fill="#555555", font=legend_font)

    current_top = top_margin
    for index, row in enumerate(rows):
        row_top = current_top
        row_height = row_heights[index]
        row_center = row_top + 30
        draw.line((left_margin - 10, row_top + row_height - 8, width - right_margin + 40, row_top + row_height - 8), fill="#f0f0f0", width=1)
        draw.text((32, row_center - 14), row.group_name, fill="black", font=label_font)

        if row.delta_attested > 0:
            color = "#2e8b57"
        elif row.delta_attested < 0:
            color = "#c0392b"
        else:
            color = "#8c8c8c"

        bar_half = abs(row.delta_attested) * scale
        if row.delta_attested == 0:
            draw.rounded_rectangle((center_x - 8, row_center - 10, center_x + 8, row_center + 10), radius=6, fill=color)
        elif row.delta_attested > 0:
            draw.rounded_rectangle((center_x, row_center - 18, center_x + bar_half, row_center + 18), radius=10, fill=color)
        else:
            draw.rounded_rectangle((center_x - bar_half, row_center - 18, center_x, row_center + 18), radius=10, fill=color)

        value_text = (
            f"№{previous_attestation}: {row.previous_attested}   "
            f"№{current_attestation}: {row.current_attested}\n"
            f"Δ {format_delta_with_percent(row.delta_attested, row.delta_attended_pct)}"
        )
        draw.multiline_text(
            (width - right_margin + 20, row_center - 24),
            value_text,
            fill="#222222",
            font=value_font,
            spacing=4,
        )

        if row.critical_note_lines:
            note_y = row_top + 56
            for note_line in row.critical_note_lines:
                draw.text((left_margin - 4, note_y), note_line, fill="#6b1f18", font=note_font)
                note_y += note_line_height

        current_top += row_height

    image_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(image_path)


def describe_attestation_comparison(comparison: AttestationComparison) -> str:
    total_delta = comparison.current_total_attested - comparison.previous_total_attested
    total_delta_pct = calc_percent_change(
        comparison.previous_total_attested,
        comparison.current_total_attested,
    )
    if total_delta > 0:
        trend = f"выросло на {format_magnitude_with_percent(total_delta, total_delta_pct)}"
    elif total_delta < 0:
        trend = f"снизилось на {format_magnitude_with_percent(total_delta, total_delta_pct)}"
    else:
        trend = "не изменилось"

    positive_rows = [row for row in comparison.rows if row.delta_attested > 0]
    negative_rows = [row for row in comparison.rows if row.delta_attested < 0]
    details = []

    if positive_rows:
        growth_group = max(positive_rows, key=lambda item: item.delta_attested)
        details.append(
            f"наибольший рост у группы {growth_group.group_name}: "
            f"{format_delta_with_percent(growth_group.delta_attested, growth_group.delta_attended_pct)}"
        )
    if negative_rows:
        decline_group = min(negative_rows, key=lambda item: item.delta_attested)
        details.append(
            f"наибольшее снижение у группы {decline_group.group_name}: "
            f"{format_magnitude_with_percent(decline_group.delta_attested, decline_group.delta_attended_pct)}"
        )

    details_text = "; ".join(details)
    if details_text:
        details_text = " " + details_text + "."

    return (
        f"По сравнению с аттестацией №{comparison.previous_attestation} общее количество аттестованных записей "
        f"по всем группам {trend}: {comparison.previous_total_attested} -> {comparison.current_total_attested}.{details_text}"
    )


def discover_image_files(input_dir: Path) -> list[Path]:
    return sorted(
        [path for path in input_dir.iterdir() if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS],
        key=lambda path: path.name.lower(),
    )


def discover_general_images(input_dir: Path) -> list[Path]:
    images = []
    for image_path in discover_image_files(input_dir):
        lower_name = image_path.name.lower()
        if lower_name.startswith("by_stud_") or lower_name.startswith("by_sub_"):
            continue
        images.append(image_path)

    def sort_key(path: Path) -> tuple[int, str]:
        stem = path.stem.lower()
        if stem in GENERAL_IMAGE_CAPTIONS:
            return (0, list(GENERAL_IMAGE_CAPTIONS).index(stem))
        return (1, path.name.lower())

    return sorted(images, key=sort_key)


def build_general_image_caption(path: Path) -> str:
    if path.stem in GENERAL_IMAGE_CAPTIONS:
        return GENERAL_IMAGE_CAPTIONS[path.stem]
    label = path.stem.replace("_", " ").strip()
    return f"Инфографика: {label}."


def discover_group_images(input_dir: Path, group_name: str) -> list[tuple[Path, str]]:
    matches: list[tuple[int, str, Path, str]] = []
    for image_path in discover_image_files(input_dir):
        if group_name not in image_path.stem:
            continue
        lower_name = image_path.name.lower()
        caption = None
        order = 100
        for prefix, builder in GROUP_IMAGE_CAPTION_BUILDERS.items():
            if lower_name.startswith(prefix):
                caption = builder(group_name)
                order = list(GROUP_IMAGE_CAPTION_BUILDERS).index(prefix)
                break
        if caption is None:
            caption = f"Инфографика по группе {group_name}: {image_path.stem.replace('_', ' ')}."
        matches.append((order, image_path.name.lower(), image_path, caption))

    matches.sort(key=lambda item: (item[0], item[1]))
    return [(path, caption) for _, _, path, caption in matches]


def find_text_report(input_dir: Path) -> Path | None:
    for file_name in ("attestation_text_report.txt", "attestation_text_report.md"):
        candidate = input_dir / file_name
        if candidate.exists():
            return candidate
    return None


def render_text_report(document: Document, text_report_path: Path | None) -> None:
    if text_report_path is None:
        return

    add_text(
        document,
        "1.4 Автоматически сформированный текстовый анализ",
        style="Heading 2",
        first_line_indent_mm=0,
    )

    text = text_report_path.read_text(encoding="utf-8-sig").strip()
    if not text:
        add_text(document, "Файл текстового анализа найден, но не содержит данных.")
        return

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        compact_line = re.sub(r"\s+", " ", line)
        if compact_line.startswith("•"):
            add_bullet(document, compact_line[1:].strip())
        elif re.match(r"^\d+(\.\d+)*[.)]?\s", compact_line):
            add_text(document, compact_line, first_line_indent_mm=0)
        elif compact_line.isupper() and len(compact_line) > 3:
            add_text(document, compact_line.title(), bold=True, first_line_indent_mm=0)
        else:
            add_text(document, compact_line)


def degree_genitive(source_dir: Path) -> str:
    return "магистратуры" if source_dir.name.lower() == "м" else "бакалавриата"


def add_title_page(document: Document, source_dir: Path, groups: list[GroupReport]) -> None:
    students_total = sum(len(group.students) for group in groups)
    records_total = sum(group.total_records for group in groups)

    add_text(
        document,
        "ОТЧЕТ",
        style="Title",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=16,
        first_line_indent_mm=0,
    )
    add_text(
        document,
        f"по аттестациям учебных групп {degree_genitive(source_dir)}",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=14,
        first_line_indent_mm=0,
    )
    add_text(document, "", first_line_indent_mm=0)
    add_text(
        document,
        f"Источник данных: {source_dir}",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent_mm=0,
    )
    add_text(
        document,
        f"Количество групп: {len(groups)}; количество студентов: {students_total}; "
        f"количество аттестационных записей: {records_total}.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent_mm=0,
    )
    add_text(
        document,
        f"Дата формирования отчета: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent_mm=0,
    )


def add_overview_section(
    document: Document,
    input_dir: Path,
    groups: list[GroupReport],
    comparison: AttestationComparison | None,
    figure_number: int,
) -> int:
    document.add_page_break()
    add_text(document, "1 Общая сводка", style="Heading 1", first_line_indent_mm=0)

    total_students = sum(len(group.students) for group in groups)
    total_records = sum(group.total_records for group in groups)
    total_failed = sum(group.total_failed for group in groups)
    total_attested = total_records - total_failed
    overall_failure_rate = total_failed / total_records * 100 if total_records else 0.0

    add_text(
        document,
        f"В отчет включены данные по {len(groups)} учебным группам {degree_genitive(input_dir)}. "
        f"Общее количество студентов составляет {total_students}, количество аттестационных "
        f"записей — {total_records}. Аттестовано {total_attested} записей, не аттестовано "
        f"{total_failed}, что соответствует доле {format_percent(overall_failure_rate)}.",
    )
    if comparison is not None:
        add_text(document, describe_attestation_comparison(comparison))

    add_text(document, "1.1 Сводные показатели по группам", style="Heading 2", first_line_indent_mm=0)
    add_table(
        document,
        [
            "Группа",
            "Студентов",
            "Дисциплин",
            "Записей",
            "Аттестовано",
            "Не аттестовано",
            "Доля неатт.",
        ],
        overview_rows(groups),
    )

    add_text(document, "1.2 Ключевые выводы", style="Heading 2", first_line_indent_mm=0)
    for group in top_problem_groups(groups):
        add_bullet(
            document,
            f"Проблемная группа: {group.name} — доля неаттестованных {format_percent(group.failure_rate)} "
            f"({group.total_failed} из {group.total_records}).",
        )
    for group in top_success_groups(groups):
        add_bullet(
            document,
            f"Наиболее стабильная группа: {group.name} — доля неаттестованных {format_percent(group.failure_rate)} "
            f"({group.total_failed} из {group.total_records}).",
        )
    for issue in overall_critical_subjects(groups):
        add_bullet(document, f"Критическая дисциплина: {issue}.")

    add_text(document, "1.3 Общая инфографика", style="Heading 2", first_line_indent_mm=0)
    for image_path in discover_general_images(input_dir):
        figure_number = add_figure(document, image_path, build_general_image_caption(image_path), figure_number)

    render_text_report(document, find_text_report(input_dir))
    return figure_number


def add_group_section(
    document: Document,
    input_dir: Path,
    group: GroupReport,
    section_number: int,
    figure_number: int,
) -> int:
    document.add_page_break()
    add_text(
        document,
        f"{section_number} Группа {group.name}",
        style="Heading 1",
        first_line_indent_mm=0,
    )

    add_text(
        document,
        f"По данным отчета «{group.report_title}» {group.execution_time}. "
        f"В группе обучается {len(group.students)} студентов. По {len(group.subject_stats)} дисциплинам "
        f"зафиксировано {group.total_records} аттестационных записей: аттестовано {group.total_attested}, "
        f"не аттестовано {group.total_failed}, доля неаттестованных составляет {format_percent(group.failure_rate)}.",
    )

    add_text(document, f"{section_number}.1 Общие показатели", style="Heading 2", first_line_indent_mm=0)
    add_table(
        document,
        ["Показатель", "Значение"],
        [
            ["Количество студентов", str(len(group.students))],
            ["Количество дисциплин", str(len(group.subject_stats))],
            ["Всего аттестационных записей", str(group.total_records)],
            ["Аттестовано", str(group.total_attested)],
            ["Не аттестовано", str(group.total_failed)],
            ["Доля неаттестованных", format_percent(group.failure_rate)],
        ],
    )

    add_text(document, f"{section_number}.2 Проблемные дисциплины", style="Heading 2", first_line_indent_mm=0)
    problem_rows = [
        [
            subject.name,
            str(subject.failed_students),
            str(subject.total_students),
            format_percent(subject.failure_rate),
        ]
        for subject in group.subject_stats[: min(7, len(group.subject_stats))]
    ]
    add_table(
        document,
        ["Дисциплина", "Не аттест.", "Студентов", "Доля неатт."],
        problem_rows,
    )

    add_text(document, f"{section_number}.3 Успешные дисциплины", style="Heading 2", first_line_indent_mm=0)
    success_rows = [
        [
            subject.name,
            str(subject.attested_students),
            str(subject.total_students),
            format_percent(subject.failure_rate),
        ]
        for subject in top_success_subjects(group)
    ]
    add_table(
        document,
        ["Дисциплина", "Аттест.", "Студентов", "Доля неатт."],
        success_rows,
    )

    add_text(
        document,
        f"{section_number}.4 Студенты с высокой неуспеваемостью >85%",
        style="Heading 2",
        first_line_indent_mm=0,
    )
    students = high_failure_students(group)
    risk_rows = (
        [
            [
                student.full_name,
                student.student_number,
                str(student.failed_count),
                format_percent(student.failure_rate),
            ]
            for student in students
        ]
        if students
        else [["Нет студентов", "-", "-", "-"]]
    )
    add_table(
        document,
        ["ФИО", "Номер", "Не аттест.", "Доля неусп."],
        risk_rows,
    )

    add_text(document, f"{section_number}.5 Инфографика", style="Heading 2", first_line_indent_mm=0)
    for image_path, caption in discover_group_images(input_dir, group.name):
        figure_number = add_figure(document, image_path, caption, figure_number)

    return figure_number


def add_teachers_section(
    document: Document,
    input_dir: Path,
    report: kafedra_teachers.KafedraReport,
    department: kafedra_teachers.Department,
    section_number: int,
    figure_number: int,
) -> int:
    document.add_page_break()
    add_text(
        document,
        f"{section_number} Аттестация по преподавателям кафедры {department.name}",
        style="Heading 1",
        first_line_indent_mm=0,
    )
    add_text(
        document,
        f"По данным отчета «Процент аттестаций по кафедрам» по состоянию на {report.execution_time}. "
        f"Учтены все дисциплины, которые ведут преподаватели кафедры, в том числе в группах "
        f"других направлений. Преподавателей — {len(department.teachers)}, групп — "
        f"{len(department.groups)}, аттестационных записей — {department.total}: аттестовано "
        f"{department.attested}, не аттестовано {department.failed}, доля неаттестованных "
        f"составляет {format_percent(department.failure_rate)}.",
    )

    add_text(document, f"{section_number}.1 Место кафедры среди кафедр вуза", style="Heading 2", first_line_indent_mm=0)
    place = report.rank_of(department.name)
    ranking = report.ranking()
    if place is not None:
        add_text(
            document,
            f"По доле неаттестованных записей кафедра {department.name} занимает "
            f"{place}-е место из {len(ranking)} (1-е место — наименьшая доля "
            f"неаттестованных): {format_percent(department.failure_rate)} при среднем по вузу "
            f"{format_percent(report.failure_rate)}.",
        )
    image_path = input_dir.parent / "kafedra_ranking.png"
    kafedra_teachers.render_ranking_chart(report, image_path, highlight=department.name)
    figure_number = add_figure(
        document,
        image_path,
        f"Доля неаттестованных записей по кафедрам вуза; кафедра {department.name} выделена.",
        figure_number,
    )

    add_text(document, f"{section_number}.2 Показатели преподавателей", style="Heading 2", first_line_indent_mm=0)
    teachers = sorted(department.teachers, key=lambda item: (-item.failure_rate, item.name))
    add_table(
        document,
        ["Преподаватель", "Групп", "Дисциплин", "Записей", "Не аттест.", "Доля неатт."],
        [
            [
                teacher.name,
                str(len(teacher.groups)),
                str(len(teacher.subjects)),
                str(teacher.total),
                str(teacher.failed),
                format_percent(teacher.failure_rate),
            ]
            for teacher in teachers
        ],
    )

    add_text(
        document,
        f"{section_number}.3 Дисциплины с долей неаттестованных не менее {FAILURE_HIGHLIGHT:.0f}%",
        style="Heading 2",
        first_line_indent_mm=0,
    )
    critical = sorted(
        (
            (teacher, record)
            for teacher, record in department.records
            if record.total and record.failure_rate >= FAILURE_HIGHLIGHT
        ),
        key=lambda item: (-item[1].failure_rate, -item[1].failed, item[0].name, item[1].group),
    )
    add_table(
        document,
        ["Преподаватель", "Группа", "Дисциплина", "Не аттест.", "Доля неатт."],
        [
            [
                teacher.name,
                record.group,
                record.subject,
                f"{record.failed} из {record.total}",
                format_percent(record.failure_rate),
            ]
            for teacher, record in critical
        ]
        or [["Нет дисциплин", "-", "-", "-", "-"]],
    )
    return figure_number


def add_risk_index_methodology(document: Document, section_number: int) -> None:
    document.add_page_break()
    add_text(
        document,
        f"{section_number} Алгоритм расчета индекса риска группы",
        style="Heading 1",
        first_line_indent_mm=0,
    )
    add_text(
        document,
        "Индекс риска — интегральная метрика от 0 до 1, где 1 соответствует максимальному риску.",
    )
    add_text(
        document,
        "risk_index = 0.5 * fail_rate~ + 0.3 * zero_rate~ + 0.2 * avg_deficit~",
        bold=True,
        first_line_indent_mm=0,
    )
    add_text(
        document,
        "Показатель рассчитывается как взвешенная сумма трех нормализованных компонентов.",
    )
    add_table(
        document,
        ["Компонент", "Вес", "Описание"],
        [
            [
                "fail_rate~",
                "0.5",
                "Доля студентов, набравших меньше минимального порога хотя бы по одной дисциплине.",
            ],
            [
                "zero_rate~",
                "0.3",
                "Доля нулевых баллов как признак полного неучастия или неявки.",
            ],
            [
                "avg_deficit~",
                "0.2",
                "Среднее недобирание до минимального порога, рассчитанное только по отрицательным отклонениям.",
            ],
        ],
    )
    add_text(
        document,
        "Перед взвешиванием к каждому компоненту применяется нормализация Min-Max scaling. "
        "Значение каждой группы приводится к диапазону от 0 до 1 относительно остальных групп выборки. "
        "Это позволяет корректно сравнивать группы независимо от абсолютных значений метрик.",
    )
    add_text(
        document,
        "x~ = (x - x_min) / (x_max - x_min)",
        bold=True,
        first_line_indent_mm=0,
    )
    add_text(
        document,
        "Здесь x~ — нормализованное значение соответствующего компонента.",
    )


def build_report(input_dir: Path, output_path: Path) -> Path:
    groups = parse_all_groups(input_dir)
    groups.sort(key=lambda item: item.name)
    comparison = build_attestation_comparison(input_dir, groups)

    document = Document()
    configure_document(document)
    add_title_page(document, input_dir, groups)

    figure_number = 1
    figure_number = add_overview_section(document, input_dir, groups, comparison, figure_number)

    next_section_number = 2
    for group in groups:
        figure_number = add_group_section(
            document,
            input_dir,
            group,
            next_section_number,
            figure_number,
        )
        next_section_number += 1

    kafedra_path = kafedra_teachers.find_kafedra_report(input_dir)
    # В отчёте по кафедрам только группы бакалавриата, поэтому раздел только в отчёте бакалавриата.
    if kafedra_path is not None and degree_genitive(input_dir) == "бакалавриата":
        kafedra_report = kafedra_teachers.parse_kafedra_report(kafedra_path)
        department = kafedra_report.department(load_department(input_dir.parent).name)
        if department is not None:
            figure_number = add_teachers_section(
                document, input_dir, kafedra_report, department, next_section_number, figure_number
            )
            next_section_number += 1

    add_risk_index_methodology(document, next_section_number)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def main() -> int:
    args = parse_args()
    input_dir = Path(args.input_dir).resolve() if args.input_dir else prompt_for_input_dir()
    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = input_dir / output_path
    output_path = output_path.resolve()

    try:
        build_report(input_dir, output_path)
    except Exception as exc:  # pragma: no cover
        print(f"Ошибка сборки отчета: {exc}", file=sys.stderr)
        return 1

    print("REPORT_CREATED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
