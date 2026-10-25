"""The short report (справка) for one level of study, with every student's scores in appendix А."""

from __future__ import annotations

import docx
from docx.table import Table

import build_short_report as short
import report_model
import synthetic


def report(tmp_path, level="б"):
    attestation_dir = synthetic.write_attestation(tmp_path)
    return docx.Document(str(short.build_short_report(attestation_dir, level)))


def text_of(document) -> str:
    cells = (cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join([*(paragraph.text for paragraph in document.paragraphs), *cells])


def appendix_tables(document) -> dict[str, Table]:
    """Group name -> its table in appendix А, matched by the caption right above the table."""
    tables, group = {}, None
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = "".join(node.text or "" for node in child.iter() if node.tag.endswith("}t"))
            if text.startswith("Таблица А."):
                group = text.split("Группа ")[-1]
        elif child.tag.endswith("}tbl") and group:
            tables[group] = Table(child, document._body)
            group = None
    return tables


def student_row(document, group: str, name: str):
    return next(row for row in appendix_tables(document)[group].rows if row.cells[1].text.startswith(name))


def fill_of(cell) -> str | None:
    shading = cell._tc.tcPr.find(docx.oxml.ns.qn("w:shd"))
    return None if shading is None else shading.get(docx.oxml.ns.qn("w:fill"))


def test_cannot_mix_master_group_into_bachelor_report(tmp_path):
    assert "ИТСм(о)-11" not in text_of(report(tmp_path)), "Bachelor short report shows a master's group"


def test_cannot_give_a_group_other_than_one_student_table(tmp_path):
    groups = list(appendix_tables(report(tmp_path)))
    assert groups == ["ИТС(о)-11", "ИСТ(о)-21"], f"Appendix must have exactly one table per group: {groups}"


def test_cannot_title_master_report_as_bachelor(tmp_path):
    assert "студентов магистратуры" in text_of(report(tmp_path, "м")), "Master short report is not titled for masters"


def test_cannot_rank_department_in_master_report(tmp_path):
    assert "-е место из" not in text_of(report(tmp_path, "м")), "Master report ranks the department by bachelor data"


def test_cannot_lose_department_rank_in_bachelor_report(tmp_path):
    assert "-е место из" in text_of(report(tmp_path)), "Bachelor report lost the department rank"


def test_cannot_lose_scores_of_a_student(tmp_path):
    row = student_row(report(tmp_path), "ИТС(о)-11", "Студент2")
    scores = sorted(cell.text for cell in row.cells[2:-1])
    assert scores == ["0", "13", "15", "2"], f"Appendix row does not match the export: {scores}"


def test_cannot_miscount_failures_of_a_student(tmp_path):
    row = student_row(report(tmp_path), "ИТС(о)-11", "Студент2")
    assert row.cells[-1].text == "2", f"Student has two scores below 13, the table says {row.cells[-1].text}"


def test_cannot_leave_zero_score_unmarked(tmp_path):
    row = student_row(report(tmp_path), "ИТС(о)-11", "Студент2")
    zero = next(cell for cell in row.cells[2:-1] if cell.text == "0")
    assert fill_of(zero) == short.SCORE_FILLS["zero"], "Zero score is not highlighted"


def test_cannot_mark_a_passing_score(tmp_path):
    row = student_row(report(tmp_path), "ИТС(о)-11", "Студент2")
    passed = next(cell for cell in row.cells[2:-1] if cell.text == "15")
    assert fill_of(passed) is None, "Score above the minimum is highlighted as a problem"


def test_cannot_print_score_too_wide_for_its_column():
    assert short.format_score(94.91) == "94,9", "Score keeps two decimals and wraps in its column"


def test_cannot_keep_dash_as_an_initial():
    assert report_model.short_name("Тестов Иван -") == "Тестов\u00a0И.", "A dash in place of the patronymic became an initial"


def test_cannot_merge_namesakes_into_one_label():
    students = [
        report_model.StudentSummary("Г", "Тестов Тест Иванович", "1", 1, 0),
        report_model.StudentSummary("Г", "Тестов Тимур Игоревич", "2", 1, 0),
    ]
    labels = short.student_labels(students)
    assert len(set(labels.values())) == 2, f"Two students got the same label: {labels}"
