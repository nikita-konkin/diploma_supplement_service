"""PDF of the DOCX reports: the best picture quality that still fits in 2 MB."""

import os

import docx
import pytest
from docx.shared import Mm
from PIL import Image

import pdf_export

needs_libreoffice = pytest.mark.skipif(pdf_export.find_soffice() is None, reason="LibreOffice is not installed")


def fake_converter(sizes: list[int], calls: list[int]):
    """Stands in for LibreOffice: ladder step -> PDF of the given size."""
    def convert(soffice, source, out_dir, profile, dpi, quality):
        step = pdf_export.LADDER.index((dpi, quality))
        calls.append(step)
        out_dir.mkdir(parents=True, exist_ok=True)
        pdf = out_dir / f"{source.stem}.pdf"
        pdf.write_bytes(b"%PDF-" + b"0" * (sizes[step] - 5))
        return pdf
    return convert


SHRINKING = [5_000_000, 3_000_000, 2_500_000, 2_100_000, 1_900_000, 1_700_000, 1_500_000, 1_200_000, 900_000]


def export(tmp_path, sizes, calls=None):
    source = tmp_path / "report.docx"
    source.write_bytes(b"PK")
    return pdf_export.export_pdf(source, soffice="soffice", converter=fake_converter(sizes, calls if calls is not None else []))


def test_cannot_pick_worse_pictures_than_what_fits(tmp_path):
    result = export(tmp_path, SHRINKING)
    assert (result.dpi, result.quality) == pdf_export.LADDER[4], f"Expected the first step under 2 MB, got {result}"


def test_cannot_return_file_over_limit_when_a_step_fits(tmp_path):
    assert export(tmp_path, SHRINKING).size <= pdf_export.MAX_BYTES, "PDF is over 2 MB although a smaller step fits"


def test_cannot_convert_twice_when_first_step_fits(tmp_path):
    calls = []
    export(tmp_path, [100_000] * len(pdf_export.LADDER), calls)
    assert calls == [0], f"A small document was converted more than once: {calls}"


def test_cannot_hide_that_nothing_fits(tmp_path):
    assert not export(tmp_path, [3_000_000] * len(pdf_export.LADDER)).fits, "PDF over 2 MB is reported as fitting"


def test_cannot_keep_a_bigger_variant_when_nothing_fits(tmp_path):
    result = export(tmp_path, [9_000_000 - step for step in range(len(pdf_export.LADDER))])
    assert (result.dpi, result.quality) == pdf_export.LADDER[-1], f"Not the smallest variant was kept: {result}"


def test_cannot_export_without_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr(pdf_export, "find_soffice", lambda: None)
    source = tmp_path / "report.docx"
    source.write_bytes(b"PK")
    with pytest.raises(pdf_export.PdfError):
        pdf_export.export_pdf(source)


@pytest.fixture(scope="module")
def heavy_pdf(tmp_path_factory):
    """Pictures of pure noise: the worst case for compression, several megabytes each."""
    folder = tmp_path_factory.mktemp("heavy")
    document = docx.Document()
    for index in range(3):
        picture = folder / f"noise{index}.png"
        Image.frombytes("RGB", (1500, 1000), os.urandom(1500 * 1000 * 3)).save(picture)
        document.add_paragraph(f"Диаграмма {index + 1}")
        document.add_picture(str(picture), width=Mm(160))
    source = folder / "heavy.docx"
    document.save(str(source))
    return pdf_export.export_pdf(source)


@needs_libreoffice
def test_cannot_exceed_two_megabytes_with_heavy_pictures(heavy_pdf):
    assert heavy_pdf.size <= pdf_export.MAX_BYTES, f"PDF with heavy pictures is {heavy_pdf.size} bytes"


@needs_libreoffice
def test_cannot_write_something_other_than_pdf(heavy_pdf):
    assert heavy_pdf.path.read_bytes()[:5] == b"%PDF-", "Export did not produce a PDF file"


@needs_libreoffice
def test_cannot_bloat_pdf_with_accessibility_tags(heavy_pdf):
    assert b"/StructTreeRoot" not in heavy_pdf.path.read_bytes(), "PDF is tagged: tables grow it several times over"
