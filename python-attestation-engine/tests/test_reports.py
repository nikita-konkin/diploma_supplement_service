"""The whole report pipeline on made-up portal exports, through the service."""

import io

import docx
import make_report
import ritm_fetch
import synthetic
from department_info import DepartmentInfo
from helpers import FakePortal, finished, log_in, make_service


def offline_reports(request, workdir, log):
    synthetic.write_attestation(workdir)
    department = DepartmentInfo(request.department, request.faculty_title, request.head_name or "И. О. Фамилия")
    period = ritm_fetch.AttestationPeriod(2026, True, 1)
    return make_report.build_all(None, period, workdir, ["б", "м"], log=log, department=department)


def finished_job(tmp_path, **overrides):
    _, client, _ = make_service(tmp_path, FakePortal(build=offline_reports))
    log_in(client)
    return client, finished(client, **overrides)


def short_report_text(client, job) -> str:
    short = next(item for item in job["downloads"] if "Краткий" in item["label"])
    document = docx.Document(io.BytesIO(client.get(short["href"]).content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )


def test_cannot_lose_any_report_file(tmp_path):
    _, job = finished_job(tmp_path)
    assert len(job["downloads"]) == 4, f"Expected short, two detailed and Markdown reports: {job}"


def test_cannot_drop_master_group_from_markdown(tmp_path):
    _, job = finished_job(tmp_path)
    assert "ИТСм(о)-11" in job["markdown"], "Master's group is missing from the Markdown report"


def test_cannot_download_short_report_as_something_else(tmp_path):
    client, job = finished_job(tmp_path)
    short = next(item for item in job["downloads"] if "Краткий" in item["label"])
    assert client.get(short["href"]).content[:2] == b"PK", "Short report download is not a DOCX file"


def test_cannot_show_department_rank_wrong(tmp_path):
    _, job = finished_job(tmp_path)
    assert "место кафедры: **2 из 3**" in job["markdown"], "Department rank in the summary is wrong"


def test_cannot_print_empty_critical_table(tmp_path):
    _, job = finished_job(tmp_path)
    assert "Таких дисциплин нет." in job["markdown"], "Empty critical-disciplines table was printed as a bare header"


def test_cannot_title_other_department_as_itrs(tmp_path):
    _, job = finished_job(tmp_path, department="ВМ", faculty_title="Факультет информатики")
    assert job["markdown"].startswith("# Кафедра ВМ:"), "Report of another department is titled ИТРС"


def test_cannot_rank_other_department_as_itrs(tmp_path):
    _, job = finished_job(tmp_path, department="ВМ", faculty_title="Факультет информатики")
    assert "место кафедры: **3 из 3**" in job["markdown"], "Rank of another department came from ИТРС"


def test_cannot_put_wrong_faculty_in_short_report_header(tmp_path):
    client, job = finished_job(tmp_path, department="ВМ", faculty_title="Факультет информатики")
    assert "Факультет информатики" in short_report_text(client, job), "Short report header has a wrong faculty"


def test_cannot_sign_short_report_by_wrong_head(tmp_path):
    client, job = finished_job(tmp_path, head_name="П. П. Петров")
    assert "П. П. Петров" in short_report_text(client, job), "Short report is signed by a wrong head"
