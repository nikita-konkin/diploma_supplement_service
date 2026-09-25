"""Synthetic workbooks in the layouts the pivot service reads."""

import io

import openpyxl
import pandas as pd

TITLES = ("наименование предмета", "часы учр", "зачет", "экзамен", "курсовой")


def statement(*students, titles=TITLES) -> bytes:
    """
    «Деканат» statement, a sheet per student: name in E1, titles on row 7.

    Args:
        students: (name, rows) pairs; a row is (subject, hours, зачет, экзамен, курсовой)
    """
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for number, (name, rows) in enumerate(students or [("Тестов Т. Т.", [("Математика", 108, None, 5, None)])]):
        sheet = workbook.create_sheet(f"Лист{number + 1}")
        sheet.append(["", "", "", "", name])
        for _ in range(5):
            sheet.append(["-"])
        sheet.append(list(titles))
        for row in rows:
            sheet.append(list(row))
    return save(workbook)


def plan(*names, column="Обязательная часть") -> bytes:
    """Discipline list: a block title, then the names."""
    content = io.BytesIO()
    pd.DataFrame({column: ["Блок 1", *names]}).to_excel(content, index=False)
    return content.getvalue()


def curriculum(credits, titles=("Структура ОП", "Объем частей ОП\nв зачетных единицах")) -> bytes:
    """Curriculum with a header, a sub-header «Всего» and a row per discipline."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["№", titles[0], "Кафедра", titles[1]])
    sheet.append(["", "", "", "Всего"])
    for number, (name, value) in enumerate(credits.items(), 1):
        sheet.append([f"Б.1.1.{number}", name, "Кафедра", value])
    return save(workbook)


def save(workbook) -> bytes:
    content = io.BytesIO()
    workbook.save(content)
    return content.getvalue()
